from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

BASE = Path(__file__).resolve().parent
OUT = BASE / "Informe_Evaluacion_Backend_Ignacio_Salinas_Diego_Perez.docx"
LOGO_PNG = BASE / "logo_informe.png"
BLUE = "1F4E79"
LIGHT = "EAF2F8"
GRAY = "D9D9D9"

logo = Image.new("RGB", (240, 240), "#55d6be")
draw = ImageDraw.Draw(logo)
draw.rounded_rectangle((0, 0, 239, 239), radius=55, fill="#55d6be")
font = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 150)
draw.text((62, 30), "N", fill="#13213c", font=font)
logo.save(LOGO_PNG)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_borders(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths=None):
    table.autofit = False
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Inches(widths[j])
            if i == 0:
                set_cell_fill(cell, BLUE)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif i % 2 == 0:
                set_cell_fill(cell, LIGHT)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_code(doc, code):
    p = doc.add_paragraph(style="Código")
    p.add_run(code)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.12
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(7)

code_style = styles.add_style("Código", WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = "Consolas"
code_style.font.size = Pt(8.5)
code_style.paragraph_format.left_indent = Inches(0.25)
code_style.paragraph_format.right_indent = Inches(0.25)
code_style.paragraph_format.space_before = Pt(5)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.0

add_page_number(section.footer.paragraphs[0])

# Portada
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.space_after = Pt(16)
p.add_run().add_picture(str(LOGO_PNG), width=Inches(1.25))
p = doc.add_paragraph("Desarrollo de una tienda online con Django", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
p = doc.add_paragraph("Informe de Evaluación Sumativa 1 de Programación Backend")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.size = Pt(14)
p.runs[0].font.color.rgb = RGBColor(80, 80, 80)
doc.add_paragraph("")
meta = doc.add_table(rows=5, cols=2)
meta_data = [
    ("Estudiantes", "Ignacio Salinas y Diego Pérez"),
    ("Asignatura", "Programación Backend"),
    ("Docente", "Claudio Rubilar Cid"),
    ("Sede", "Apoquindo"),
    ("Fecha", "7 de septiembre de 2026"),
]
for i, (k, v) in enumerate(meta_data):
    meta.cell(i, 0).text = k
    meta.cell(i, 1).text = v
    set_cell_borders(meta.cell(i, 0)); set_cell_borders(meta.cell(i, 1))
    set_cell_margins(meta.cell(i, 0)); set_cell_margins(meta.cell(i, 1))
    set_cell_fill(meta.cell(i, 0), LIGHT)
    meta.cell(i, 0).paragraphs[0].runs[0].font.bold = True
meta.columns[0].width = Inches(1.7)
meta.columns[1].width = Inches(4.9)
doc.add_page_break()

doc.add_heading("Resumen del proyecto", level=1)
doc.add_paragraph(
    "Desarrollamos Nexo Store, una aplicación web básica que promociona productos tecnológicos. "
    "La solución separa la lógica en dos aplicaciones Django: productos administra el catálogo y sus detalles, "
    "mientras usuarios presenta el perfil del equipo. El resultado cumple los requerimientos de navegación, "
    "colecciones de datos, plantillas, archivos estáticos, identidad visual y uso de un framework de interfaz."
)
doc.add_heading("Objetivo", level=1)
doc.add_paragraph(
    "Implementar una tienda online funcional con Django que permita recorrer una página de inicio, consultar un "
    "catálogo de cuatro productos, abrir el detalle de cada uno y conocer al equipo responsable del proyecto."
)
doc.add_heading("Tecnologías utilizadas", level=1)
tech = doc.add_table(rows=1, cols=3)
for j, h in enumerate(("Tecnología", "Versión o uso", "Propósito")): tech.cell(0, j).text = h
for row in [
    ("Python", "3.13", "Lógica del servidor"),
    ("Django", "5.2.6", "Vistas, rutas y plantillas"),
    ("Bootstrap", "5.3.3", "Interfaz adaptable"),
    ("HTML y CSS", "HTML5 y CSS3", "Estructura e identidad visual"),
    ("Git", "Control de versiones", "Historial del proyecto"),
]:
    cells = tech.add_row().cells
    for j, value in enumerate(row): cells[j].text = value
format_table(tech, [1.45, 1.65, 3.55])

doc.add_heading("Cumplimiento de la pauta", level=1)
req = doc.add_table(rows=1, cols=3)
for j, h in enumerate(("Requerimiento", "Implementación", "Evidencia")): req.cell(0, j).text = h
for row in [
    ("Dos aplicaciones", "productos y usuarios", "INSTALLED_APPS y rutas independientes"),
    ("Tres rutas de menú", "Inicio, Catálogo y Nosotros", "Barra de navegación compartida"),
    ("Cuatro productos", "Catálogo tecnológico", "Cuatro tarjetas enlazadas a detalles"),
    ("Diccionarios", "PRODUCTOS en data.py", "Cada producto contiene id, nombre, categoría, precio, descripción e imagen"),
    ("Templates y static", "Plantillas por app y carpeta global", "HTML, CSS y cinco imágenes SVG"),
    ("Identidad y framework", "Nexo Store y Bootstrap", "Logotipo, colores, tipografía y diseño responsive"),
]:
    cells = req.add_row().cells
    for j, value in enumerate(row): cells[j].text = value
format_table(req, [1.45, 2.0, 3.2])

doc.add_heading("Arquitectura de la solución", level=1)
doc.add_paragraph(
    "El proyecto sigue la estructura Modelo Vista Plantilla de Django. En esta primera versión no se necesita una "
    "base de datos para el catálogo, porque la pauta solicita organizar la información en colecciones de datos. "
    "Por eso utilizamos un diccionario central y lo entregamos a las plantillas mediante el contexto de cada vista."
)
add_code(doc, """Proyecto_Backend_Salinas_Perez/
├── tienda_backend/       configuración y rutas principales
├── productos/            datos, vistas, rutas y plantillas del catálogo
├── usuarios/             vista, ruta y plantilla del equipo
├── templates/base.html   navegación y estructura compartida
├── static/css/           estilos propios
├── static/images/        logotipo e ilustraciones
├── manage.py
└── requirements.txt""")
doc.add_heading("Flujo de una solicitud", level=2)
doc.add_paragraph(
    "Cuando una persona entra a una URL, Django revisa la configuración de rutas. La ruta selecciona una vista; "
    "la vista consulta el diccionario de productos, prepara un contexto y renderiza la plantilla correspondiente. "
    "La plantilla hereda de base.html para mantener el mismo menú y pie de página en todo el sitio."
)

doc.add_heading("Implementación del catálogo", level=1)
doc.add_paragraph(
    "Los productos se almacenan en un diccionario cuya clave es un número entero. Esta decisión permite recuperar "
    "un producto por su identificador y demuestra el uso de variables, operadores, listas por comprensión y estructuras de decisión."
)
add_code(doc, """PRODUCTOS = {
    1: {
        \"id\": 1,
        \"nombre\": \"Audífonos Pulse Pro\",
        \"categoria\": \"Audio\",
        \"precio\": 59990,
        \"descripcion\": \"Audífonos inalámbricos con cancelación de ruido...\",
        \"imagen\": \"images/audifonos.svg\",
        \"destacado\": True,
    },
    # Tres productos adicionales
}""")
doc.add_heading("Vistas y decisiones", level=2)
doc.add_paragraph(
    "La vista de catálogo transforma los valores del diccionario en una lista. Si la URL incluye una categoría, "
    "se aplica un filtro; de lo contrario se muestran los cuatro elementos. La vista de detalle utiliza get para "
    "buscar el identificador y responde con un error 404 cuando el producto no existe."
)
add_code(doc, """def catalogo(request):
    categoria = request.GET.get(\"categoria\", \"todas\")
    productos = list(PRODUCTOS.values())
    if categoria != \"todas\":
        productos = [p for p in productos if p[\"categoria\"].lower() == categoria.lower()]
    return render(request, \"productos/catalogo.html\", {
        \"productos\": productos,
        \"categoria_actual\": categoria,
    })""")

doc.add_heading("Rutas y navegación", level=1)
doc.add_paragraph(
    "El archivo de rutas principal delega las direcciones a cada aplicación. productos controla el inicio, el "
    "catálogo y los detalles, mientras usuarios controla la página del equipo. El menú presenta tres accesos "
    "visibles y se reutiliza desde la plantilla base."
)
routes = doc.add_table(rows=1, cols=3)
for j, h in enumerate(("Ruta", "Vista", "Resultado")): routes.cell(0, j).text = h
for row in [
    ("/", "inicio", "Portada y productos destacados"),
    ("/catalogo/", "catalogo", "Cuatro productos y filtros"),
    ("/producto/<id>/", "detalle", "Imagen, descripción y precio"),
    ("/usuarios/perfil/", "perfil", "Información de los estudiantes"),
]:
    cells = routes.add_row().cells
    for j, value in enumerate(row): cells[j].text = value
format_table(routes, [1.65, 1.35, 3.7])

doc.add_page_break()
doc.add_heading("Resultado visual", level=1)
doc.add_paragraph(
    "La interfaz usa la identidad Nexo Store, con una paleta azul, verde y tonos claros. Bootstrap aporta la grilla "
    "adaptable, la navegación móvil y componentes consistentes. Los estilos propios definen tarjetas, portada y jerarquía visual."
)
for image, caption in [
    ("captura_inicio.png", "Figura 1 Página de inicio con navegación y productos destacados"),
    ("captura_catalogo.png", "Figura 2 Catálogo completo con cuatro productos"),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(BASE / image), width=Inches(6.35))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

doc.add_heading("Pruebas realizadas", level=1)
doc.add_paragraph(
    "Ejecutamos la verificación de configuración de Django y seis pruebas automáticas. También recorrimos la aplicación "
    "en un navegador para comprobar que las rutas, enlaces, imágenes, textos y detalles se mostraran correctamente."
)
tests = doc.add_table(rows=1, cols=2)
tests.cell(0, 0).text = "Prueba"; tests.cell(0, 1).text = "Resultado"
for name in [
    "La página de inicio responde y muestra productos destacados",
    "El catálogo contiene exactamente cuatro productos",
    "El filtro Audio devuelve dos productos",
    "El detalle válido presenta el producto y su precio",
    "Un identificador inexistente devuelve error 404",
    "La página Nosotros muestra a ambos estudiantes",
]:
    cells = tests.add_row().cells
    cells[0].text = name; cells[1].text = "Aprobada"
format_table(tests, [5.35, 1.35])
doc.add_paragraph("Resultado de ejecución: 6 pruebas aprobadas, 0 errores.")

doc.add_heading("Instrucciones de ejecución", level=1)
doc.add_paragraph("Para ejecutar la aplicación en Windows se deben usar los siguientes comandos desde la carpeta del proyecto:")
add_code(doc, """python -m venv .venv
.venv\\Scripts\\Activate.ps1
pip install -r requirements.txt
python manage.py migrate
python manage.py runserver""")
doc.add_paragraph("Luego se abre la dirección http://127.0.0.1:8000/ en el navegador.")

doc.add_heading("Conclusiones", level=1)
doc.add_paragraph(
    "El desarrollo permitió aplicar la estructura de un proyecto Django y separar responsabilidades mediante aplicaciones "
    "reutilizables. La navegación conecta correctamente las distintas vistas y el diccionario central simplifica la "
    "administración de los productos solicitados. Además, el uso conjunto de plantillas heredadas, archivos estáticos y "
    "Bootstrap produce una interfaz clara y consistente."
)
doc.add_paragraph(
    "Como mejora futura, el diccionario puede reemplazarse por modelos de base de datos para administrar productos desde "
    "el panel de Django. También se podrían incorporar registro real de clientes, carro de compras y control de inventario."
)
doc.add_heading("Referencias", level=1)
doc.add_paragraph("Django Software Foundation. Django documentation. https://docs.djangoproject.com/")
doc.add_paragraph("Bootstrap Team. Bootstrap 5 documentation. https://getbootstrap.com/docs/5.3/")
doc.add_paragraph("Python Software Foundation. Python documentation. https://docs.python.org/3/")

doc.core_properties.title = "Desarrollo de una tienda online con Django"
doc.core_properties.subject = "Evaluación Sumativa 1 de Programación Backend"
doc.core_properties.author = "Ignacio Salinas y Diego Pérez"
doc.core_properties.keywords = "Django, Python, backend, eCommerce"
doc.save(OUT)
print(OUT)
